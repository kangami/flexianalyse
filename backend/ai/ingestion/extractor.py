"""
Document Extractor using Docling.

Handles extraction of content from all supported file types :
  - PDF (with OCR for scanned pages)
  - DOCX, PPTX
  - XLSX, CSV
  - HTML, Markdown
  - Plain text, code files

Returns structured chunks ready for embedding.
"""

import hashlib
import io
import logging
import os
import tempfile
from dataclasses import dataclass, field
from datetime import datetime
from typing import Optional

from docling.document_converter import DocumentConverter, PdfFormatOption
from docling.datamodel.base_models import InputFormat
from docling.datamodel.pipeline_options import (
    PdfPipelineOptions,
    RapidOcrOptions,
    PictureDescriptionApiOptions,
)
from docling.chunking import HybridChunker
from langchain_text_splitters import RecursiveCharacterTextSplitter

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Supported MIME types
# ---------------------------------------------------------------------------

SUPPORTED_MIME_TYPES = {
    # Documents
    'application/pdf':                                                              'pdf',
    'application/vnd.google-apps.document':                                         'gdoc',
    'application/vnd.openxmlformats-officedocument.wordprocessingml.document':      'docx',
    'application/msword':                                                           'docx',

    # Spreadsheets
    'application/vnd.google-apps.spreadsheet':                                      'sheet',
    'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet':            'xlsx',
    'text/csv':                                                                     'csv',

    # Presentations
    'application/vnd.google-apps.presentation':                                     'pptx',
    'application/vnd.openxmlformats-officedocument.presentationml.presentation':    'pptx',

    # Text / Code
    'text/plain':       'text',
    'text/markdown':    'text',
    'text/html':        'html',
    'application/json': 'text',
    'text/x-python':    'text',
    'text/javascript':  'text',
}

EXCLUDED_PREFIXES = ('video/', 'audio/', 'image/gif', 'image/mp4', 'image/webp')

EXCLUDED_EXTENSIONS = {
    # Archives et binaires
    'apk', 'ipa', 'exe', 'dll', 'so', 'dylib',
    'zip', 'tar', 'gz', 'rar', '7z', 'bz2',
    'jar', 'war', 'ear', 'class',
    # Images
    'jpg', 'jpeg', 'png', 'gif', 'bmp', 'ico',
    'tiff', 'webp', 'svg', 'psd', 'ai',
    # Fonts
    'ttf', 'otf', 'woff', 'woff2', 'eot',
    # Données binaires
    'bin', 'dat', 'db', 'sqlite', 'pkl',
    # Compiled
    'pyc', 'pyo', 'o', 'a',
}

# Use semantic chunking for rich formats, fixed for flat text
SEMANTIC_FORMATS = {'pdf', 'docx', 'gdoc', 'pptx', 'html'}
FIXED_FORMATS    = {'text', 'csv', 'xlsx', 'sheet'}

FIXED_CHUNK_SIZE    = int(os.getenv('CHUNK_SIZE', '500'))
FIXED_CHUNK_OVERLAP = int(os.getenv('CHUNK_OVERLAP', '50'))

# Minimum fraction of the document's text the structure-aware chunker must
# capture; below this we rebuild from the full markdown so no content is lost.
EXTRACT_MIN_COVERAGE = float(os.getenv('EXTRACT_MIN_COVERAGE', '0.85'))


def _alnum_len(s: str) -> int:
    """Count alphanumeric characters — a formatting-agnostic measure of how much
    actual text content a string holds (ignores markdown pipes, '#', spaces)."""
    return sum(1 for c in (s or '') if c.isalnum())


# ---------------------------------------------------------------------------
# Data classes
# ---------------------------------------------------------------------------

@dataclass
class ExtractedChunk:
    content: str
    chunk_index: int
    chunk_type: str                     # 'text', 'table', 'title', 'code'
    page_number: Optional[int] = None
    section_title: Optional[str] = None
    token_count: Optional[int] = None
    chunk_metadata: dict = field(default_factory=dict)


@dataclass
class ExtractionResult:
    resource_title: str
    file_format: str
    content_hash: str
    chunks: list[ExtractedChunk]
    raw_size_bytes: int
    error: Optional[str] = None

    @property
    def success(self) -> bool:
        return self.error is None and len(self.chunks) > 0


# ---------------------------------------------------------------------------
# Extractor
# ---------------------------------------------------------------------------

class DocumentExtractor:
    """
    Extracts and chunks content from any supported document type.
    Uses Docling for rich formats (PDF/DOCX/PPTX) and simple splitters for flat text.
    """

    def __init__(self):
        # OCR (RapidOCR / ONNX Runtime) is very memory-heavy and OOMs
        # (std::bad_alloc) on large/scanned PDFs. Most business PDFs are "native"
        # (carry a text layer) and need no OCR. DOCLING_OCR controls this:
        #   "auto"  (default) — OCR only PDFs with no detectable text layer
        #   "true"            — always OCR
        #   "false"           — never OCR
        mode = os.getenv("DOCLING_OCR", "auto").strip().lower()
        if mode in ("1", "true", "yes"):
            mode = "true"
        elif mode in ("0", "false", "no"):
            mode = "false"
        elif mode != "auto":
            mode = "auto"
        self._ocr_mode = mode

        try:
            self._image_scale = float(os.getenv("DOCLING_IMAGE_SCALE", "1.0"))
        except (TypeError, ValueError):
            self._image_scale = 1.0

        # VLM image/chart description via GPT-4o-vision — OFF by default.
        # FLEXIANALYSE_VLM=true → describe figures/charts/diagrams with GPT-4o.
        self._vlm_enabled = os.getenv("FLEXIANALYSE_VLM", "false").lower() in ("1", "true", "yes")

        # Default converter: NO OCR (fast, low memory) — used for native-text
        # PDFs and every other semantic format.
        self._converter = self._build_converter(do_ocr=False)
        # OCR converter: built eagerly only when OCR is forced on, otherwise
        # lazily the first time a scanned PDF is detected (saves memory).
        self._ocr_converter = (
            self._build_converter(do_ocr=True) if self._ocr_mode == "true" else None
        )
        logger.info(
            "Docling extractor ready (OCR mode=%s, engine=PaddleOCR, VLM=%s)",
            self._ocr_mode, self._vlm_enabled,
        )

        self._hybrid_chunker = HybridChunker(
            tokenizer="BAAI/bge-small-en-v1.5",  # tokenizer for chunk sizing
            max_tokens=FIXED_CHUNK_SIZE,
            merge_peers=True,
        )

        self._fixed_splitter = RecursiveCharacterTextSplitter(
            chunk_size=FIXED_CHUNK_SIZE * 4,     # ~chars, not tokens
            chunk_overlap=FIXED_CHUNK_OVERLAP * 4,
            separators=["\n\n", "\n", " ", ""],
        )

    # =========================================================================
    # OCR helpers
    # =========================================================================

    # Sample a few pages spread across the doc; below this many chars/page on
    # average we treat the PDF as scanned (image-only) and enable OCR.
    _OCR_SAMPLE_PAGES = 5
    _OCR_MIN_CHARS_PER_PAGE = 30

    def _build_converter(self, do_ocr: bool) -> DocumentConverter:
        """Create a Docling converter with OCR on or off."""
        pdf_options = PdfPipelineOptions()
        pdf_options.do_ocr = do_ocr
        pdf_options.do_table_structure = True
        if do_ocr:
            # PaddleOCR engine = RapidOCR with the PaddlePaddle backend (avoids
            # the onnxruntime std::bad_alloc seen with the default ONNX backend).
            pdf_options.ocr_options = RapidOcrOptions(backend="paddle")
            pdf_options.images_scale = self._image_scale

        # Optional VLM (GPT-4o-vision) description of figures/charts/diagrams.
        if self._vlm_enabled:
            pdf_options.do_picture_description = True
            pdf_options.enable_remote_services = True   # required for remote API
            pdf_options.generate_picture_images = True  # render images for the VLM
            pdf_options.images_scale = max(self._image_scale, 1.0)
            pdf_options.picture_description_options = PictureDescriptionApiOptions(
                url="https://api.openai.com/v1/chat/completions",
                headers={"Authorization": f"Bearer {os.getenv('OPENAI_API_KEY', '')}"},
                params={"model": os.getenv("VLM_MODEL", "gpt-4o")},
                prompt=(
                    "Describe this image, chart, diagram or table in detail. "
                    "Include any data, numbers, axis labels, legends and trends "
                    "so it can be searched by text."
                ),
                timeout=60,
            )

        return DocumentConverter(
            format_options={
                InputFormat.PDF: PdfFormatOption(pipeline_options=pdf_options)
            }
        )

    def _get_ocr_converter(self) -> DocumentConverter:
        """Lazily build (and cache) the OCR-enabled converter."""
        if self._ocr_converter is None:
            logger.info("Building OCR converter (scanned PDF detected)")
            self._ocr_converter = self._build_converter(do_ocr=True)
        return self._ocr_converter

    def _pdf_has_text_layer(self, path: str) -> bool:
        """Heuristic: does this PDF carry an extractable text layer (native PDF)
        vs. being a scan (image-only)? Samples pages and measures text density.
        On any failure, assume native — that skips OCR and avoids the OOM."""
        try:
            import pypdfium2 as pdfium
        except Exception:
            logger.warning("pypdfium2 unavailable — skipping OCR detection")
            return True

        pdf = None
        try:
            pdf = pdfium.PdfDocument(path)
            n = len(pdf)
            if n == 0:
                return True

            idxs = sorted({
                int(i * (n - 1) / max(1, self._OCR_SAMPLE_PAGES - 1))
                for i in range(min(self._OCR_SAMPLE_PAGES, n))
            })
            total_chars = 0
            for i in idxs:
                page = pdf[i]
                textpage = page.get_textpage()
                total_chars += len((textpage.get_text_range() or "").strip())
                textpage.close()
                page.close()

            avg = total_chars / len(idxs)
            has_text = avg >= self._OCR_MIN_CHARS_PER_PAGE
            logger.info(
                "PDF text-layer check: %.0f chars/page over %d sampled pages → %s",
                avg, len(idxs), "native (no OCR)" if has_text else "scanned (OCR)",
            )
            return has_text
        except Exception as e:
            logger.warning("PDF text-layer detection failed (%s) — assuming native", e)
            return True
        finally:
            if pdf is not None:
                try:
                    pdf.close()
                except Exception:
                    pass

    # =========================================================================
    # PUBLIC
    # =========================================================================

    def is_supported(self, mime_type: str, filename: str) -> bool:
        """Cheap pre-download check: will this file type be processed at all?
        Lets the ingestion pipeline skip huge unsupported binaries (ISO, APK,
        MP4, images, archives...) BEFORE downloading them."""
        return self._detect_format(mime_type or "", filename or "") is not None

    def detect_format(self, mime_type: str, filename: str):
        """Public: resolve the internal format ('pdf', 'docx', 'csv'...) or None."""
        return self._detect_format(mime_type or "", filename or "")

    def sample_text(self, raw_content: bytes, file_format: str, filename: str,
                    max_chars: int = 3000) -> str:
        """Cheap text sample for the ingestion router (no full Docling parse).
        PDFs: first pages via pypdfium2; flat formats: decoded head. Best-effort."""
        try:
            if file_format in ('csv', 'xlsx', 'sheet', 'text', 'html'):
                return (self._decode_to_text(raw_content, file_format, filename) or "")[:max_chars]
            if file_format == 'pdf':
                import pypdfium2 as pdfium
                with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as t:
                    t.write(raw_content)
                    p = t.name
                try:
                    pdf = pdfium.PdfDocument(p)
                    parts = []
                    for i in range(min(2, len(pdf))):
                        page = pdf[i]
                        tp = page.get_textpage()
                        parts.append(tp.get_text_range() or "")
                        tp.close()
                        page.close()
                    pdf.close()
                    return ("\n".join(parts))[:max_chars]
                finally:
                    os.unlink(p)
        except Exception as e:
            logger.warning("sample_text failed for %s: %s", filename, e)
        return ""

    def extract(
        self,
        raw_content: bytes,
        mime_type: str,
        filename: str,
        strategy: str = "generic",
    ) -> ExtractionResult:
        """
        Main entry point — extract and chunk a document.

        Args:
            raw_content : raw bytes of the file
            mime_type   : MIME type string
            filename    : original filename (used for extension fallback)
            strategy    : ingestion strategy chosen by the router. Currently only
                          'generic' changes nothing; per-strategy extraction is
                          wired in the next phase. Recorded for visibility.

        Returns:
            ExtractionResult with chunks ready for embedding
        """
        if strategy and strategy != "generic":
            logger.info("Extraction strategy '%s' requested for %s (generic path for now)",
                        strategy, filename)

        file_format = self._detect_format(mime_type, filename)

        if file_format is None:
            return ExtractionResult(
                resource_title=filename,
                file_format='unsupported',
                content_hash='',
                chunks=[],
                raw_size_bytes=len(raw_content),
                error=f"Unsupported MIME type: {mime_type}"
            )

        content_hash = hashlib.sha256(raw_content).hexdigest()

        try:
            if file_format in SEMANTIC_FORMATS:
                chunks = self._extract_semantic(raw_content, file_format, filename)
            else:
                chunks = self._extract_fixed(raw_content, file_format, filename)

            if not chunks:
                return ExtractionResult(
                    resource_title=filename,
                    file_format=file_format,
                    content_hash=content_hash,
                    chunks=[],
                    raw_size_bytes=len(raw_content),
                    error="No extractable text found in document (empty or unreadable)",
                )

            return ExtractionResult(
                resource_title=filename,
                file_format=file_format,
                content_hash=content_hash,
                chunks=chunks,
                raw_size_bytes=len(raw_content),
            )

        except Exception as e:
            logger.error(f"Extraction failed for {filename}: {e}", exc_info=True)
            return ExtractionResult(
                resource_title=filename,
                file_format=file_format,
                content_hash=content_hash,
                chunks=[],
                raw_size_bytes=len(raw_content),
                error=str(e)
            )

    # =========================================================================
    # SEMANTIC EXTRACTION (PDF, DOCX, PPTX, HTML)
    # =========================================================================

    def _extract_semantic(
        self,
        raw_content: bytes,
        file_format: str,
        filename: str,
    ) -> list[ExtractedChunk]:
        """Use Docling HybridChunker for structure-aware chunking.
        Falls back to markdown export + fixed splitter when the chunker
        yields nothing (scanned PDFs, simple layouts, etc.).
        """

        # Write to temp file — Docling needs a file path
        suffix = f".{file_format}" if file_format not in ('gdoc',) else '.docx'
        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
            tmp.write(raw_content)
            tmp_path = tmp.name

        try:
            # Pick OCR vs no-OCR per PDF: always for "true", never for "false",
            # and only for image-only (scanned) PDFs in "auto" mode.
            converter = self._converter
            if file_format == 'pdf':
                if self._ocr_mode == 'true':
                    converter = self._get_ocr_converter()
                elif self._ocr_mode == 'auto' and not self._pdf_has_text_layer(tmp_path):
                    converter = self._get_ocr_converter()

            result = converter.convert(tmp_path)
            doc = result.document

            # --- Pass 1: HybridChunker (semantic, structure-aware) ---
            chunks: list[ExtractedChunk] = []
            try:
                for i, chunk in enumerate(self._hybrid_chunker.chunk(doc)):
                    text = chunk.text.strip()
                    if not text:
                        continue

                    page_num = None
                    section = None
                    chunk_type = 'text'

                    if hasattr(chunk, 'meta'):
                        meta = chunk.meta
                        if hasattr(meta, 'headings') and meta.headings:
                            section = meta.headings[-1]
                        if hasattr(meta, 'doc_items'):
                            for item in meta.doc_items:
                                # Page number lives in provenance (prov)
                                if page_num is None:
                                    prov = getattr(item, 'prov', None)
                                    if prov:
                                        page_num = getattr(prov[0], 'page_no', None)
                                label = getattr(item, 'label', '')
                                if 'table' in str(label).lower():
                                    chunk_type = 'table'
                                    break
                                elif 'title' in str(label).lower() or 'heading' in str(label).lower():
                                    chunk_type = 'title'
                                    break
                                elif 'code' in str(label).lower():
                                    chunk_type = 'code'
                                    break

                    chunks.append(ExtractedChunk(
                        content=text,
                        chunk_index=i,
                        chunk_type=chunk_type,
                        page_number=page_num,
                        section_title=section,
                        token_count=len(text.split()),
                        chunk_metadata={'filename': filename, 'format': file_format}
                    ))
            except Exception as e:
                logger.warning(f"HybridChunker failed for {filename}: {e}")

            # --- Pass 2: coverage guard against the FULL markdown export ---
            # The structure-aware chunker can silently miss content (tables,
            # multi-column layouts, captions). We compare what it captured to the
            # complete markdown representation; if it covers too little, we
            # rebuild from the markdown so NOTHING is lost (quality > structure).
            md = ""
            try:
                md = doc.export_to_markdown() or ""
            except Exception as e:
                logger.warning(f"Markdown export failed for {filename}: {e}")

            hc_alnum = sum(_alnum_len(c.content) for c in chunks)
            md_alnum = _alnum_len(md)
            coverage = (hc_alnum / md_alnum) if md_alnum else 1.0

            if md.strip() and (not chunks or coverage < EXTRACT_MIN_COVERAGE):
                logger.info(
                    f"{filename}: chunker captured {coverage:.0%} of the text "
                    f"(hc={hc_alnum} vs md={md_alnum}) — rebuilding from markdown "
                    f"to avoid content loss"
                )
                md_chunks: list[ExtractedChunk] = []
                for i, chunk_text in enumerate(self._fixed_splitter.split_text(md)):
                    chunk_text = chunk_text.strip()
                    if chunk_text:
                        md_chunks.append(ExtractedChunk(
                            content=chunk_text,
                            chunk_index=i,
                            chunk_type='text',
                            token_count=len(chunk_text.split()),
                            chunk_metadata={
                                'filename': filename,
                                'format': file_format,
                                'extraction': 'markdown_full',
                            }
                        ))
                if md_chunks:
                    chunks = md_chunks

            # VLM descriptions of figures/charts become searchable chunks.
            if self._vlm_enabled:
                chunks = chunks + self._picture_chunks(doc, filename, file_format)

            logger.info(
                f"Semantic extraction: {filename} → {len(chunks)} chunks "
                f"(coverage {coverage:.0%})"
            )
            return chunks

        finally:
            os.unlink(tmp_path)

    def _picture_chunks(self, doc, filename: str, file_format: str) -> list[ExtractedChunk]:
        """Turn Docling picture descriptions (VLM annotations) into chunks so
        charts/diagrams become searchable. Best-effort — never raises."""
        out: list[ExtractedChunk] = []
        try:
            for idx, pic in enumerate(getattr(doc, 'pictures', []) or []):
                desc = None
                for ann in (getattr(pic, 'annotations', []) or []):
                    txt = getattr(ann, 'text', None)
                    if txt and txt.strip():
                        desc = txt.strip()
                        break
                if not desc:
                    continue
                page = None
                prov = getattr(pic, 'prov', None)
                if prov:
                    page = getattr(prov[0], 'page_no', None)
                out.append(ExtractedChunk(
                    content=f"[Figure] {desc}",
                    chunk_index=100000 + idx,
                    chunk_type='figure',
                    page_number=page,
                    token_count=len(desc.split()),
                    chunk_metadata={
                        'filename': filename,
                        'format': file_format,
                        'extraction': 'vlm_picture',
                    },
                ))
            if out:
                logger.info(f"VLM described {len(out)} figure(s) in {filename}")
        except Exception as e:
            logger.warning(f"Picture description extraction failed for {filename}: {e}")
        return out

    # =========================================================================
    # FIXED EXTRACTION (CSV, XLSX, plain text, code)
    # =========================================================================

    def _extract_fixed(
        self,
        raw_content: bytes,
        file_format: str,
        filename: str,
    ) -> list[ExtractedChunk]:
        """Use RecursiveCharacterTextSplitter for flat/tabular content."""

        text = self._decode_to_text(raw_content, file_format, filename)
        if not text:
            return []

        raw_chunks = self._fixed_splitter.split_text(text)
        chunks = []
        for i, chunk_text in enumerate(raw_chunks):
            chunk_text = chunk_text.strip()
            if not chunk_text:
                continue
            chunks.append(ExtractedChunk(
                content=chunk_text,
                chunk_index=i,
                chunk_type='table' if file_format in ('csv', 'xlsx', 'sheet') else 'text',
                token_count=len(chunk_text.split()),
                chunk_metadata={'filename': filename, 'format': file_format}
            ))

        logger.info(f"Fixed extraction: {filename} → {len(chunks)} chunks")
        return chunks

    def _decode_to_text(self, raw_content: bytes, file_format: str, filename: str) -> str:
        """Convert raw bytes to plain text for flat formats."""
        try:
            if file_format == 'csv':
                return raw_content.decode('utf-8', errors='replace')

            elif file_format in ('xlsx', 'sheet'):
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(raw_content), read_only=True, data_only=True)
                lines = []
                for sheet in wb.worksheets:
                    lines.append(f"## Sheet: {sheet.title}")
                    for row in sheet.iter_rows(values_only=True):
                        row_text = '\t'.join(str(c) if c is not None else '' for c in row)
                        if row_text.strip():
                            lines.append(row_text)
                return '\n'.join(lines)

            else:
                # text, markdown, json, code
                return raw_content.decode('utf-8', errors='replace')

        except Exception as e:
            logger.error(f"Decode failed for {filename}: {e}")
            return ''

    # =========================================================================
    # HELPERS
    # =========================================================================

    def _detect_format(self, mime_type: str, filename: str) -> Optional[str]:
        """Detect file format from MIME type, fallback to extension."""

        # Check excluded types first
        for prefix in EXCLUDED_PREFIXES:
            if mime_type.startswith(prefix):
                return None

        # Check excluded extensions
        ext = filename.rsplit('.', 1)[-1].lower() if '.' in filename else ''
        if ext in EXCLUDED_EXTENSIONS:
            logger.info(f"Excluded extension: {filename}")
            return None

        # Check supported types
        if mime_type in SUPPORTED_MIME_TYPES:
            return SUPPORTED_MIME_TYPES[mime_type]

        # Fallback to file extension
        ext = filename.rsplit('.', 1)[-1].lower() if '.' in filename else ''
        ext_map = {
            'pdf': 'pdf', 'docx': 'docx', 'doc': 'docx',
            'pptx': 'pptx', 'ppt': 'pptx',
            'xlsx': 'xlsx', 'xls': 'xlsx',
            'csv': 'csv', 'txt': 'text',
            'md': 'text', 'html': 'html',
            'py': 'text', 'js': 'text', 'ts': 'text',
            'json': 'text', 'yaml': 'text', 'yml': 'text',
        }
        return ext_map.get(ext)
